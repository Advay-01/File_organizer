"""Regression tests for mislabeled text formats and parser fallbacks."""
from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app import api
from app.core import history
from app.core.content_cache import ContentIntelligenceCache
from app.core.content_intelligence import ContentSuggestion, plan_content_intelligence
from app.core.evidence import extract_evidence
from app.core.extractor import extract_content
from app.models.schemas import FileInfo


def file_info(path: Path) -> FileInfo:
    return FileInfo(
        path=str(path),
        filename=path.name,
        extension=path.suffix.lower(),
        size_kb=1,
        date_created="2026-01-01",
    )


def completion_for(
    folder: str,
    subfolder: str,
    quote: str,
    suggested_name: str,
    conflict: str = "misleading",
):
    def complete(_messages):
        return json.dumps({
            "files": [{
                "id": 0,
                "category": folder,
                "subcategory": subfolder,
                "confidence": "high",
                "reasoning": "The locally detected content format is explicit.",
                "evidence_quote": quote,
                "filename_conflict": conflict,
                "suggested_name": suggested_name,
                "name_evidence": quote,
            }]
        })

    return complete


class MisleadingExtensionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cache = ContentIntelligenceCache(self.root / "intelligence.db")

    def tearDown(self):
        self.cache.close()
        self.temp.cleanup()

    def assert_extension_corrected(
        self,
        filename: str,
        content: str,
        expected_extension: str,
        suggested_stem: str,
        quote: str,
    ):
        source = self.root / filename
        source.write_text(content, encoding="utf-8")
        original = source.read_bytes()
        model_name = f"{suggested_stem}{source.suffix}"
        result = plan_content_intelligence(
            [file_info(source)],
            str(self.root),
            ["fake/model"],
            cache=self.cache,
            completion_fn=completion_for("Detected", "Files", quote, model_name),
        )
        item = result.suggestions[0]

        self.assertEqual(item.detected_type, expected_extension)
        self.assertEqual(item.filename_conflict, "misleading")
        self.assertEqual(item.suggested_name, f"{suggested_stem}{expected_extension}")
        self.assertIn("corrected the extension", item.rename_reason)
        self.assertEqual(source.read_bytes(), original)

    def test_csv_data_mislabeled_as_html_gets_csv_extension(self):
        content = (
            "region,quarter,revenue,units\n"
            "North,Q1,120000,24\n"
            "South,Q1,98000,18\n"
            "East,Q1,151000,31\n"
        ) * 4
        self.assert_extension_corrected(
            "terms_of_service.html",
            content,
            ".csv",
            "q1_regional_sales_metrics",
            "region,quarter,revenue,units",
        )

    def test_cpp_source_mislabeled_as_json_gets_cpp_extension(self):
        content = (
            "#include <iostream>\n"
            "#include <cstring>\n"
            "using namespace std;\n"
            "int main() {\n"
            "  char name[8];\n"
            "  strcpy(name, \"dangerous overflow\");\n"
            "  cout << name << endl;\n"
            "  return 0;\n"
            "}\n"
        ) * 4
        self.assert_extension_corrected(
            "user_feedback_batch_1.json",
            content,
            ".cpp",
            "vulnerable_buffer_overflow_demo",
            "#include <iostream>",
        )

    def test_sql_mislabeled_as_python_gets_sql_extension(self):
        content = (
            "SELECT username, password FROM users WHERE id = 1;\n"
            "' OR '1'='1' --\n"
            "DROP TABLE users;\n"
        ) * 6
        self.assert_extension_corrected(
            "python_test_script.py",
            content,
            ".sql",
            "sql_injection_payload_samples",
            "SELECT username, password FROM users",
        )

    def test_html_mislabeled_as_markdown_gets_html_extension(self):
        content = (
            "<!DOCTYPE html><html><head><title>Team status</title></head>"
            "<body><script>alert('xss')</script><p>Weekly update</p></body></html>\n"
        ) * 5
        self.assert_extension_corrected(
            "weekly_team_status.md",
            content,
            ".html",
            "xss_exploit_payload",
            "<!DOCTYPE html>",
        )

    def test_identical_csv_bytes_do_not_share_extension_specific_cache_context(self):
        content = (
            "patient_id,diagnosis_code,admission_date,discharge_status\n"
            "PAT-90112,ICD10-J44.9,2026-09-12,DISCHARGED_HOME\n"
            "PAT-90113,ICD10-I10,2026-09-14,TRANSFERRED_ICU\n"
            "PAT-90114,ICD10-E11.9,2026-09-15,DISCHARGED_HOME\n"
        )
        first = self.root / "q4_financial_audit.json"
        second = self.root / "marketing_campaign_notes.docx"
        first.write_text(content, encoding="utf-8")
        second.write_text(content, encoding="utf-8")
        seen = []

        def completion(messages):
            payload = json.loads(messages[1]["content"])
            seen.append(payload["files"])
            return json.dumps({
                "files": [{
                    "id": row["id"],
                    "category": "Healthcare",
                    "subcategory": "Audit Data",
                    "confidence": "high",
                    "reasoning": "The CSV header and rows are explicit.",
                    "evidence_quote": "patient_id",
                    "filename_conflict": "misleading",
                    "suggested_name": f"{Path(row['current_name']).stem}.csv",
                    "name_evidence": "patient_id",
                } for row in payload["files"]]
            })

        result = plan_content_intelligence(
            [file_info(first), file_info(second)],
            str(self.root),
            ["fake/model"],
            cache=self.cache,
            completion_fn=completion,
        )

        prompt_rows = seen[0]
        self.assertEqual(
            sorted((row["current_name"], row["extension"]) for row in prompt_rows),
            [("marketing_campaign_notes.docx", ".docx"), ("q4_financial_audit.json", ".json")],
        )
        self.assertEqual(
            [item.suggested_name for item in result.suggestions],
            ["q4_financial_audit.csv", "marketing_campaign_notes.csv"],
        )
        self.assertEqual(first.read_text(encoding="utf-8"), content)
        self.assertEqual(second.read_text(encoding="utf-8"), content)


class ParserFallbackTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_bash_text_in_xml_uses_text_fallback(self):
        source = self.root / "nightly_backup.xml"
        source.write_text(
            "#!/bin/bash\nset -e\ntar -czf backup.tar.gz /srv/data\n",
            encoding="utf-8",
        )
        extraction = extract_content(str(source))
        self.assertEqual(extraction.parser_status, "fallback")
        self.assertEqual(extraction.detected_type, ".sh")
        self.assertEqual(extraction.detection_confidence, "high")
        self.assertIn("#!/bin/bash", extraction.text)
        self.assertIn("XML", extraction.warning)

    def test_plain_text_in_xlsx_remains_readable_text(self):
        source = self.root / "release_notes.xlsx"
        source.write_text(
            "Release checklist\nConfirm the installer and tests before shipping.\n",
            encoding="utf-8",
        )
        extraction = extract_content(str(source))
        self.assertEqual(extraction.parser_status, "fallback")
        self.assertEqual(extraction.detected_type, ".txt")
        self.assertEqual(extraction.detection_confidence, "high")
        self.assertIn("Release checklist", extraction.text)
        self.assertIn("XLSX", extraction.warning)

    def test_valid_structured_documents_still_parse(self):
        json_path = self.root / "valid.json"
        json_path.write_text('{"topic": "database migration", "status": "ready"}', encoding="utf-8")
        xml_path = self.root / "valid.xml"
        xml_path.write_text(
            '<?xml version="1.0" encoding="utf-8"?>'
            "<notes><topic>database migration</topic><status>ready</status></notes>",
            encoding="utf-8",
        )
        xlsx_path = self.root / "valid.xlsx"
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Region", "Revenue"])
        sheet.append(["North", 120000])
        workbook.save(xlsx_path)
        docx_path = self.root / "valid.docx"
        import docx

        document = docx.Document()
        document.add_paragraph("database migration project status ready")
        document.save(docx_path)
        png_path = self.root / "valid.png"
        from PIL import Image

        Image.new("RGB", (8, 8), "blue").save(png_path)

        json_result = extract_content(str(json_path))
        xml_result = extract_content(str(xml_path))
        xlsx_result = extract_content(str(xlsx_path))
        docx_result = extract_content(str(docx_path))
        png_result = extract_content(str(png_path))

        self.assertEqual((json_result.parser_status, json_result.detected_type), ("parsed", ".json"))
        self.assertEqual((xml_result.parser_status, xml_result.detected_type), ("parsed", ".xml"))
        self.assertEqual((xlsx_result.parser_status, xlsx_result.detected_type), ("parsed", ".xlsx"))
        self.assertEqual((docx_result.parser_status, docx_result.detected_type), ("parsed", ".docx"))
        self.assertEqual((png_result.parser_status, png_result.detected_type), ("parsed", ".png"))
        self.assertIn("database migration", json_result.text)
        self.assertIn("database migration", xml_result.text)
        self.assertIn("North | 120000", xlsx_result.text)
        self.assertIn("database migration project status ready", docx_result.text)
        self.assertIn("Image: PNG; 8x8", png_result.text)

    def test_binary_content_in_structured_extension_remains_unavailable(self):
        source = self.root / "corrupt.xlsx"
        source.write_bytes(b"\x00\x01PK\x03\x04\xff\x00" * 20)
        extraction = extract_content(str(source))
        self.assertEqual(extraction.parser_status, "unavailable")
        self.assertEqual(extraction.text, "")

        calls = []
        cache = ContentIntelligenceCache(self.root / "cache.db")
        try:
            result = plan_content_intelligence(
                [file_info(source)],
                str(self.root),
                ["fake/model"],
                cache=cache,
                completion_fn=lambda _: calls.append(1) or "{}",
            )
        finally:
            cache.close()
        self.assertEqual(calls, [])
        self.assertEqual(result.suggestions[0].suggested_name, source.name)
        self.assertEqual(result.suggestions[0].extraction_status, "unavailable")
        self.assertEqual(source.read_bytes(), b"\x00\x01PK\x03\x04\xff\x00" * 20)

    def test_malformed_docx_and_png_are_not_decoded_as_text(self):
        samples = {
            "broken.docx": b"PK\x03\x04" + b"\x00\xff\x01\x02" * 40,
            "broken.png": b"\x89PNG\r\n\x1a\n" + b"\x00\xff\x01\x02" * 40,
        }
        for name, payload in samples.items():
            with self.subTest(name=name):
                source = self.root / name
                source.write_bytes(payload)
                extraction = extract_content(str(source))
                self.assertEqual(extraction.text, "")
                self.assertEqual(extraction.parser_status, "unavailable")

    def test_missing_pdf_crypto_provider_uses_fallback_instead_of_aborting(self):
        source = self.root / "protected.pdf"
        source.write_bytes(b"%PDF-1.7\nencrypted fixture\n%%EOF")
        from pypdf.errors import DependencyError

        with patch(
            "app.core.extractor._read_pdf",
            side_effect=DependencyError("cryptography>=3.1 is required for AES algorithm"),
        ):
            extraction = extract_content(str(source), max_chars=200)

        self.assertIn(extraction.parser_status, {"fallback", "unavailable"})
        self.assertIn("PDF", extraction.warning)

    def test_aes_encrypted_pdf_does_not_abort_extraction(self):
        source = self.root / "aes-protected.pdf"
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("", owner_password="sortly-test-owner", algorithm="AES-256")
        with source.open("wb") as stream:
            writer.write(stream)

        extraction = extract_content(str(source), max_chars=200)

        self.assertEqual(extraction.detected_type, ".pdf")
        self.assertIn(extraction.parser_status, {"parsed", "fallback", "unavailable"})

    def test_uncertain_text_format_keeps_original_extension(self):
        source = self.root / "meeting_notes.md"
        source.write_text(
            "The team discussed the current project, upcoming work, and ownership. "
            "No structured format or source language is present in this note. " * 4,
            encoding="utf-8",
        )
        cache = ContentIntelligenceCache(self.root / "cache.db")
        try:
            result = plan_content_intelligence(
                [file_info(source)],
                str(self.root),
                ["fake/model"],
                cache=cache,
                completion_fn=completion_for(
                    "Work",
                    "Meeting Notes",
                    "The team discussed the current project",
                    "meeting_summary.txt",
                    conflict="none",
                ),
            )
        finally:
            cache.close()
        self.assertEqual(result.suggestions[0].suggested_name, "meeting_summary.md")
        self.assertFalse(extract_evidence(file_info(source)).extension_mismatch)
        self.assertTrue(source.exists())

    def test_chess_record_keeps_pgn_extension_and_is_not_a_config_file(self):
        source = self.root / "chess_com_games_2026-10-17.pgn"
        source.write_text(
            '[Event "Casual Game"]\n[Site "Chess.com"]\n[Date "2026.10.17"]\n'
            '[White "Alice"]\n[Black "Bob"]\n[Result "1-0"]\n\n'
            '1. e4 e5 2. Nf3 Nc6 3. Bb5 a6 1-0\n', encoding="utf-8",
        )
        evidence = extract_evidence(file_info(source))
        self.assertEqual(evidence.detected_type, ".pgn")
        self.assertFalse(evidence.extension_mismatch)
        self.assertFalse(any("conflicts with detected" in warning for warning in evidence.warnings))

        cache = ContentIntelligenceCache(self.root / "chess-cache.db")
        try:
            result = plan_content_intelligence(
                [file_info(source)], str(self.root), ["fake/model"], cache=cache,
                completion_fn=completion_for(
                    "Documents", "Chess Games", "Casual Game", "chess_games.txt",
                    conflict="extension_mismatch",
                ),
            )
        finally:
            cache.close()
        item = result.suggestions[0]
        self.assertEqual(item.filename_conflict, "none")
        self.assertTrue(item.suggested_name.endswith(".pgn"))

    def test_patch_and_diff_keep_their_extensions(self):
        content = (
            "diff --git a/settings.ini b/settings.ini\n"
            "--- a/settings.ini\n+++ b/settings.ini\n"
            "@@ -1,2 +1,2 @@\n"
            " [settings]\n-old=true\n+old=false\n"
        )
        for extension in (".patch", ".diff"):
            with self.subTest(extension=extension):
                source = self.root / f"fix-settings{extension}"
                source.write_text(content, encoding="utf-8")
                evidence = extract_evidence(file_info(source))
                self.assertEqual(evidence.detected_type, ".patch")
                self.assertFalse(evidence.extension_mismatch)

    def test_generic_readable_text_does_not_prove_an_unknown_extension_is_wrong(self):
        source = self.root / "field_notes.custom"
        source.write_text("Unstructured field observations about plants and soil. " * 4, encoding="utf-8")
        evidence = extract_evidence(file_info(source))
        self.assertEqual(evidence.detected_type, "unknown")
        self.assertFalse(evidence.extension_mismatch)
        self.assertEqual(evidence.extraction_status, "sufficient")

    def test_sensitive_fallback_is_not_sent_to_cloud(self):
        source = self.root / "private.xml"
        source.write_text(
            "-----BEGIN RSA PRIVATE KEY-----\n"
            "not-a-real-key-material\n"
            "-----END RSA PRIVATE KEY-----\n",
            encoding="utf-8",
        )
        cache = ContentIntelligenceCache(self.root / "cache.db")
        calls = []
        try:
            result = plan_content_intelligence(
                [file_info(source)],
                str(self.root),
                ["fake/cloud-model"],
                cache=cache,
                completion_fn=lambda _: calls.append(1) or "{}",
            )
        finally:
            cache.close()
        self.assertEqual(calls, [])
        item = result.suggestions[0]
        self.assertTrue(item.sensitive)
        self.assertFalse(item.selected)
        self.assertNotIn("BEGIN RSA PRIVATE KEY", item.evidence_excerpt)
        self.assertEqual(source.name, item.suggested_name)

    def test_text_disguised_as_png_is_detected_as_sensitive_env(self):
        source = self.root / "hero_banner_background.png"
        source.write_text(
            "# Synthetic credentials for tests\n"
            "AWS_ACCESS_KEY_ID=AKIAABCDEFGHIJKLMNOP\n"
            "AWS_SECRET_ACCESS_KEY=not-a-real-secret-value\n"
            "OPENAI_API_KEY=sk-abcdefghijklmnopqrstuvwx\n",
            encoding="utf-8",
        )
        extraction = extract_content(str(source))
        self.assertEqual(extraction.parser_status, "fallback")
        self.assertEqual(extraction.detected_type, ".env")
        self.assertEqual(extraction.detection_confidence, "high")

        cache = ContentIntelligenceCache(self.root / "cache.db")
        calls = []
        try:
            result = plan_content_intelligence(
                [file_info(source)],
                str(self.root),
                ["fake/cloud-model"],
                cache=cache,
                completion_fn=lambda _: calls.append(1) or "{}",
            )
        finally:
            cache.close()
        self.assertEqual(calls, [])
        item = result.suggestions[0]
        self.assertTrue(item.sensitive)
        self.assertFalse(item.selected)
        self.assertEqual(item.target_path, "Misc/Sensitive Review")
        self.assertNotEqual(item.target_path, "Misc/Unreadable")
        self.assertNotIn("not-a-real-secret-value", item.evidence_excerpt)

    def test_explicit_redacted_cloud_processing_removes_env_secrets(self):
        source = self.root / "hero_banner_background.png"
        source.write_text(
            "AWS_SECRET_ACCESS_KEY=not-a-real-secret-value\n"
            "OPENAI_API_KEY=sk-abcdefghijklmnopqrstuvwx\n",
            encoding="utf-8",
        )
        sent = []
        cache = ContentIntelligenceCache(self.root / "cache.db")

        def completion(messages):
            sent.append(messages[1]["content"])
            return json.dumps({"files": [{
                "id": 0,
                "category": "Security",
                "subcategory": "Credentials",
                "confidence": "high",
                "reasoning": "Sensitive environment data was explicitly approved for redacted review.",
                "evidence_quote": "AWS_SECRET_ACCESS_KEY",
                "filename_conflict": "misleading",
                "suggested_name": "hero_banner_background.env",
                "name_evidence": "AWS_SECRET_ACCESS_KEY",
            }]})

        try:
            result = plan_content_intelligence(
                [file_info(source)],
                str(self.root),
                ["fake/cloud-model"],
                cache=cache,
                completion_fn=completion,
                allow_sensitive_cloud=True,
            )
        finally:
            cache.close()
        self.assertEqual(len(sent), 1)
        self.assertNotIn("not-a-real-secret-value", sent[0])
        self.assertNotIn("sk-abcdefghijklmnopqrstuvwx", sent[0])
        self.assertIn("[REDACTED SECRET]", sent[0])
        self.assertTrue(result.suggestions[0].sensitive)


class ExtensionApplyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.saved_home = os.environ.get("ORGANIZER_HOME")
        os.environ["ORGANIZER_HOME"] = str(self.root / "home")
        api.plans.clear()

    def tearDown(self):
        api.plans.clear()
        if self.saved_home is None:
            os.environ.pop("ORGANIZER_HOME", None)
        else:
            os.environ["ORGANIZER_HOME"] = self.saved_home
        self.temp.cleanup()

    def test_corrected_extension_collision_and_undo(self):
        first = self.root / "first.html"
        second = self.root / "second.html"
        first.write_text("region,revenue\nNorth,120000\n", encoding="utf-8")
        second.write_text("region,revenue\nSouth,98000\n", encoding="utf-8")
        suggestions = [
            ContentSuggestion(
                id=index,
                path=str(path),
                current_name=path.name,
                category="Data",
                subcategory="Sales",
                confidence="high",
                reasoning="CSV data.",
                evidence_quote="region,revenue",
                name_evidence="region,revenue",
                evidence_excerpt="region,revenue",
                evidence_sources=["raw_text_fallback"],
                warnings=[],
                filename_conflict="misleading",
                suggested_name="q1_sales.csv",
                rename_reason="Detected CSV content and corrected the extension.",
                sufficient_for_change=True,
                detected_type=".csv",
                extraction_status="sufficient",
            )
            for index, path in enumerate((first, second))
        ]
        moves = api._content_moves(suggestions, self.root)
        self.assertEqual(Path(moves[0].dest).name, "q1_sales.csv")
        self.assertEqual(Path(moves[1].dest).name, "q1_sales-1.csv")
        api.plans["extension-plan"] = api.StoredPlan(self.root, moves, suggestions)

        result = api.apply(api.ApplyRequest(
            plan_id="extension-plan",
            files=[
                api.ApplyFile(
                    id=str(index),
                    accepted=True,
                    target_path="Data/Sales",
                    destination_name=Path(move.dest).name,
                )
                for index, move in enumerate(moves)
            ],
        ))

        self.assertEqual(result["moved"], 2)
        self.assertTrue((self.root / "Data" / "Sales" / "q1_sales.csv").exists())
        self.assertTrue((self.root / "Data" / "Sales" / "q1_sales-1.csv").exists())
        self.assertFalse(first.exists())
        self.assertFalse(second.exists())

        undo = history.undo_last_run(str(self.root))
        self.assertEqual(undo["restored"], 2)
        self.assertTrue(first.exists())
        self.assertTrue(second.exists())

    def test_preview_exposes_corrected_extension_without_moving_file(self):
        source = self.root / "terms_of_service.html"
        source.write_text("region,revenue\nNorth,120000\n", encoding="utf-8")
        original = source.read_bytes()
        suggestion = ContentSuggestion(
            id=0,
            path=str(source),
            current_name=source.name,
            category="Business",
            subcategory="Sales",
            confidence="high",
            reasoning="CSV data.",
            evidence_quote="region,revenue",
            name_evidence="region,revenue",
            evidence_excerpt="region,revenue",
            evidence_sources=["text_format_detection"],
            warnings=[],
            filename_conflict="misleading",
            suggested_name="q1_regional_sales_metrics.csv",
            rename_reason="Detected CSV content and corrected the extension.",
            sufficient_for_change=True,
            detected_type=".csv",
            extraction_status="sufficient",
        )
        scan_report = SimpleNamespace(files_seen=1, skipped=0, reasons={}, truncated=False)
        with (
            patch("app.api.scan_directory_with_report", return_value=([file_info(source)], scan_report)),
            patch("app.api.provider_store.activate_runtime", return_value=["fake/model"]),
            patch(
                "app.api.plan_content_intelligence",
                return_value=SimpleNamespace(suggestions=[suggestion], stats={}),
            ),
        ):
            response = api.preview(api.PreviewRequest(source=str(self.root), rename=True))

        self.assertEqual(response.files[0].name, source.name)
        self.assertEqual(response.files[0].destination_name, "q1_regional_sales_metrics.csv")
        self.assertTrue(response.files[0].renamed)
        self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
