"""Structured local evidence extraction for content-first organization."""
from __future__ import annotations

import json
import re
import shutil
import zipfile
from dataclasses import asdict, dataclass, field
from html.parser import HTMLParser
from pathlib import Path
import subprocess
import tempfile
from typing import List
from xml.etree import ElementTree

from app.core.extractor import _read_text, extract_content, extract_snippet
from app.core.privacy import detect_sensitive, detect_suspicious_commands, redact_sensitive
from app.models.schemas import FileInfo
from app.core.safety import snapshot_file

EVIDENCE_VERSION = 6
_XMP_LIMIT = 256 * 1024
_XMP_MM = "http://ns.adobe.com/xap/1.0/mm/"
_XMP_REF = "http://ns.adobe.com/xap/1.0/sType/ResourceRef#"


@dataclass
class EvidenceBundle:
    path: str
    fingerprint: str
    filename: str
    extension: str
    parent_directory: str
    content: str = ""
    metadata: dict = field(default_factory=dict)
    sources: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    detected_type: str = "unknown"
    detected_type_confidence: str = "unknown"
    extension_mismatch: bool = False
    parser_status: str = "unavailable"
    extraction_status: str = "unavailable"
    ocr_used: bool = False
    sensitive_findings: List[dict] = field(default_factory=list)
    redacted_content: str = ""
    suspicious_commands: List[str] = field(default_factory=list)

    def evidence_text(self) -> str:
        metadata_text = " ".join(
            str(value) for value in self.metadata.values() if value not in (None, "")
        )
        return f"{self.content} {metadata_text}".strip()

    def redacted_evidence_text(self) -> str:
        metadata_text = " ".join(
            redact_sensitive(str(value))
            for value in self.metadata.values()
            if value not in (None, "")
        )
        return f"{self.redacted_content or redact_sensitive(self.content)} {metadata_text}".strip()

    @property
    def sensitive(self) -> bool:
        return bool(self.sensitive_findings)

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False)

    @classmethod
    def from_json(cls, value: str) -> "EvidenceBundle":
        return cls(**json.loads(value))


def content_fingerprint(path: Path) -> str | None:
    """Hash the complete file so same-size mid-file edits cannot hit stale cache data."""
    try:
        return snapshot_file(path).sha256
    except OSError:
        return None


def _zip_type(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        if "[Content_Types].xml" in names:
            if any(name.startswith("word/") for name in names):
                return ".docx"
            if any(name.startswith("ppt/") for name in names):
                return ".pptx"
            if any(name.startswith("xl/") for name in names):
                return ".xlsx"
        if "mimetype" in names:
            return ".epub"
        return ".zip"
    except (OSError, ValueError, zipfile.BadZipFile):
        return ".zip"


def detect_file_type(path: Path) -> str:
    """Detect common formats from signatures without trusting the extension."""
    try:
        with path.open("rb") as source:
            header = source.read(32)
    except OSError:
        return "unknown"
    if header.startswith(b"%PDF"):
        return ".pdf"
    if header.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if header.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if header.startswith((b"GIF87a", b"GIF89a")):
        return ".gif"
    if header.startswith(b"BM"):
        return ".bmp"
    if header.startswith((b"II*\x00", b"MM\x00*")):
        return ".tiff"
    if header.startswith(b"RIFF") and header[8:12] == b"WEBP":
        return ".webp"
    if header.startswith(b"PK\x03\x04"):
        return _zip_type(path)
    if header.startswith(b"\x1f\x8b"):
        return ".gz"
    if header.startswith(b"Rar!\x1a\x07"):
        return ".rar"
    return "unknown"


def _extensions_match(extension: str, detected: str) -> bool:
    if detected == "unknown" or not extension:
        return True
    aliases = {
        ".jpeg": ".jpg",
        ".tif": ".tiff",
        ".xlsm": ".xlsx",
        ".tgz": ".tar",
        ".htm": ".html",
        ".diff": ".patch",
    }
    return aliases.get(extension, extension) == aliases.get(detected, detected)


def _clean_metadata(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:500]


def _xmp_relationship_metadata(raw: object) -> dict:
    """Read bounded, embedded identifiers without resolving any resource reference."""
    if not isinstance(raw, (str, bytes)) or len(raw) > _XMP_LIMIT:
        return {}
    if isinstance(raw, str):
        raw = raw.encode("utf-8")
    if len(raw) > _XMP_LIMIT:
        return {}
    if b"\x00" in raw or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        return {}
    try:
        root = ElementTree.fromstring(raw)
    except (ValueError, LookupError, ElementTree.ParseError):
        return {}
    elements = list(root.iter())
    if len(elements) > 2048:
        return {}
    values: dict[str, set[str]] = {}

    def remember(key: str, value: object) -> None:
        if isinstance(value, str):
            value = value.strip()
            if 3 <= len(value) <= 256 and not re.search(r"[\s\x00-\x1f<>\"']", value):
                values.setdefault(key, set()).add(value)

    for local, key in (("DocumentID", "xmp_document_id"), ("InstanceID", "xmp_instance_id"),
                       ("OriginalDocumentID", "xmp_original_document_id")):
        tag = f"{{{_XMP_MM}}}{local}"
        for element in elements:
            if element.tag == tag:
                remember(key, element.text)
            remember(key, element.attrib.get(tag))
    for derived in root.iter(f"{{{_XMP_MM}}}DerivedFrom"):
        for element in derived.iter():
            for local, key in (("documentID", "xmp_derived_from_document_id"),
                               ("instanceID", "xmp_derived_from_instance_id")):
                tag = f"{{{_XMP_REF}}}{local}"
                if element.tag == tag:
                    remember(key, element.text)
                remember(key, element.attrib.get(tag))
    if any(len(items) > 1 for items in values.values()):
        return {"xmp_relationship_conflict": "Embedded document identifiers contradict each other."}
    return {key: next(iter(items)) for key, items in values.items()}


class _LocalResourceParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.references: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        fields = dict(attrs)
        reference = fields.get("src") if tag in {"img", "script", "source", "video", "audio"} else None
        if tag == "link" and "stylesheet" in (fields.get("rel") or "").casefold().split():
            reference = fields.get("href")
        if reference and len(reference) <= 2048 and len(self.references) < 64:
            if reference not in self.references:
                self.references.append(reference)


def _has_semantic_metadata(metadata: dict) -> bool:
    """Dimensions, provenance IDs, software names and authors do not prove a subject."""
    generic = re.compile(
        r"^(?:untitled|document|presentation|slide\s*show|microsoft\s+(?:word|powerpoint|excel))"
        r"(?:[\s_-]*\d+)?(?:\.(?:docx?|pptx?|pdf|xlsx?))?$", re.I,
    )
    for key in ("title", "subject", "keywords", "description", "headings"):
        value = metadata.get(key)
        if not isinstance(value, str):
            continue
        value = value.strip()
        if len(value) >= 8 and re.search(r"[^\W\d_]", value, re.UNICODE) and not generic.fullmatch(value):
            return True
    return False


def _document_metadata(path: Path, extension: str) -> dict:
    metadata = {}
    try:
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                metadata["warning"] = "encrypted PDF metadata may be incomplete"
            raw = reader.metadata or {}
            for source, target in (
                ("/Title", "title"),
                ("/Author", "author"),
                ("/Subject", "subject"),
                ("/Creator", "creator"),
            ):
                value = _clean_metadata(raw.get(source))
                if value:
                    metadata[target] = value
            xmp = reader.root_object.get("/Metadata")
            if xmp:
                metadata.update(_xmp_relationship_metadata(xmp.get_object().get_data()))
        elif extension == ".docx":
            import docx

            props = docx.Document(str(path)).core_properties
            for source, target in (
                ("title", "title"),
                ("author", "author"),
                ("subject", "subject"),
                ("keywords", "keywords"),
            ):
                value = _clean_metadata(getattr(props, source, ""))
                if value:
                    metadata[target] = value
        elif extension == ".pptx":
            from pptx import Presentation

            props = Presentation(str(path)).core_properties
            for source, target in (("title", "title"), ("author", "author"), ("subject", "subject")):
                value = _clean_metadata(getattr(props, source, ""))
                if value:
                    metadata[target] = value
        elif extension in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                props = workbook.properties
                for source, target in (("title", "title"), ("creator", "author"), ("subject", "subject")):
                    value = _clean_metadata(getattr(props, source, ""))
                    if value:
                        metadata[target] = value
            finally:
                workbook.close()
        elif extension in {".jpg", ".jpeg", ".png", ".heic", ".webp", ".tiff", ".tif", ".bmp", ".gif"}:
            from PIL import Image

            if extension == ".heic":
                import pillow_heif

                pillow_heif.register_heif_opener()
            with Image.open(path) as image:
                exif = image.getexif()
                for tag, target in ((270, "description"), (315, "artist"), (33432, "copyright")):
                    value = _clean_metadata(exif.get(tag))
                    if value:
                        metadata[target] = value
                metadata["dimensions"] = f"{image.width}x{image.height}"
                metadata["image_format"] = image.format or ""
                xmp = image.info.get("xmp") or image.info.get("XML:com.adobe.xmp") or exif.get(700)
                metadata.update(_xmp_relationship_metadata(xmp))
        elif extension in {".html", ".htm"}:
            with path.open("r", encoding="utf-8", errors="ignore") as source:
                raw = source.read(20000)
            title = re.search(r"<title[^>]*>(.*?)</title>", raw, re.I | re.S)
            if title:
                metadata["title"] = _clean_metadata(re.sub(r"<[^>]+>", " ", title.group(1)))
            headings = re.findall(r"<h[1-3][^>]*>(.*?)</h[1-3]>", raw, re.I | re.S)
            if headings:
                metadata["headings"] = " | ".join(
                    _clean_metadata(re.sub(r"<[^>]+>", " ", value)) for value in headings[:5]
                )
            resources = _LocalResourceParser()
            resources.feed(raw)
            if resources.references:
                metadata["local_dependencies"] = resources.references
        elif extension == ".xmp":
            with path.open("rb") as source:
                metadata.update(_xmp_relationship_metadata(source.read(_XMP_LIMIT + 1)))
    except Exception:
        metadata["warning"] = "metadata could not be read"
    return metadata


def _ocr_image(path: Path, max_chars: int) -> tuple[str, bool]:
    executable = shutil.which("tesseract")
    if not executable:
        return "", False
    try:
        from PIL import Image

        with tempfile.TemporaryDirectory(prefix="sortly-ocr-") as temp:
            raster = Path(temp) / "page.png"
            with Image.open(path) as image:
                image.thumbnail((1600, 1600))
                image.convert("RGB").save(raster)
            completed = subprocess.run(
                [executable, str(raster), "stdout", "--psm", "3"],
                stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                text=True, encoding="utf-8", errors="replace", check=True, timeout=10,
            )
            text = completed.stdout
        return re.sub(r"\s+", " ", text).strip()[:max_chars], bool(text.strip())
    except Exception:
        return "", False


def _ocr_pdf(path: Path, max_chars: int) -> tuple[str, bool]:
    """Inspect at most two low-resolution pages with already installed local tools."""
    renderer = shutil.which("pdftoppm")
    if not renderer or not shutil.which("tesseract"):
        return "", False
    try:
        with tempfile.TemporaryDirectory(prefix="sortly-pdf-ocr-") as temp:
            prefix = Path(temp) / "page"
            subprocess.run(
                [renderer, "-f", "1", "-l", "2", "-r", "120", "-scale-to", "1600",
                 "-png", str(path), str(prefix)],
                stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                check=True, timeout=15,
            )
            parts = []
            for raster in sorted(Path(temp).glob("page-*.png"))[:2]:
                if raster.stat().st_size > 12 * 1024 * 1024:
                    continue
                text, _used = _ocr_image(raster, max_chars)
                if text:
                    parts.append(text)
                if sum(map(len, parts)) >= max_chars:
                    break
            content = " ".join(parts)[:max_chars]
            return content, bool(content)
    except Exception:
        return "", False


def extract_evidence(
    file_info: FileInfo,
    max_chars: int = 4000,
    fingerprint: str | None = None,
) -> EvidenceBundle:
    path = Path(file_info.path)
    fingerprint = fingerprint if fingerprint is not None else (content_fingerprint(path) or "")
    extension = path.suffix.lower()
    signature_type = detect_file_type(path)
    extraction = extract_content(str(path), max_chars=max_chars)
    detected_type = signature_type
    detection_confidence = "high" if signature_type != "unknown" else extraction.detection_confidence
    if signature_type == "unknown" and extraction.detected_type != "unknown":
        detected_type = extraction.detected_type

    warnings = []
    if extraction.warning:
        warnings.append(extraction.warning)
    # Readable prose is not evidence that an unfamiliar text extension is wrong.
    # A failed structured parser is different: its fallback can establish that
    # the declared format (for example XLSX) is actually plain text.
    if detected_type == ".txt" and extraction.parser_status == "sniffed_text":
        detected_type = "unknown"
        detection_confidence = "unknown"
    if detected_type != "unknown" and not _extensions_match(extension, detected_type):
        warnings.append(
            f"Extension {extension or '(none)'} conflicts with detected format {detected_type}."
        )

    content = extraction.text
    if re.fullmatch(r"Image:\s*[A-Za-z0-9]+\s*;\s*\d+x\d+", content.strip()):
        # The image parser's format/dimension summary is not semantic content.
        content = ""
    metadata = _document_metadata(path, extension)
    if metadata.pop("warning", None):
        warnings.append("Metadata could not be fully read; the file may be encrypted or malformed.")
    if metadata.get("warning"):
        warnings.append(str(metadata["warning"]))

    ocr_used = False
    if len(content.strip()) < 20 and signature_type == ".pdf":
        ocr_content, ocr_used = _ocr_pdf(path, max_chars)
        if ocr_content:
            content = ocr_content
            warnings.append("Text came from optional PDF OCR of up to two pages and should be reviewed.")
        else:
            warnings.append("No PDF text was extracted. Optional local Poppler and Tesseract OCR may be required; review this file.")
    if len(content.strip()) < 20 and extension in {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"}:
        content, ocr_used = _ocr_image(path, max_chars)
        if content:
            warnings.append("Text came from optional OCR and should be reviewed.")
        elif extension:
            warnings.append("No OCR runtime is available for this image.")

    sources = []
    if content:
        if extraction.parser_status == "fallback":
            sources.append("raw_text_fallback")
        else:
            sources.append("extracted_content" if not ocr_used else "ocr")
    if extraction.parser_status == "fallback":
        sources.append("parser_fallback")
    if metadata:
        sources.append("metadata")
    if signature_type != "unknown":
        sources.append("file_signature")
    elif extraction.detected_type != "unknown":
        sources.append("text_format_detection")
    semantic_metadata = _has_semantic_metadata(metadata)
    status = "sufficient" if len(content.strip()) >= 80 else (
        "weak" if content.strip() or semantic_metadata else "unavailable"
    )
    if status == "unavailable":
        warnings.append("No reliable content or metadata could be extracted.")
    # PDF object offsets, font tables, and other container bytes are not
    # document content. Their numeric IDs can accidentally pass a card checksum.
    # Keep the larger raw-text scan for genuine text files (including .env).
    sensitivity_text = content if signature_type != "unknown" else (_read_text(path, max_chars) or content)
    findings = detect_sensitive(
        sensitivity_text,
        [
            *[str(value) for value in metadata.values() if value not in (None, "")],
            path.name,
            path.parent.name,
        ],
    )
    if findings:
        kinds = ", ".join(sorted({finding.kind.replace("_", " ") for finding in findings}))
        warnings.append(
            f"Potential sensitive content detected ({kinds}). Cloud processing is disabled by default."
        )

    suspicious = detect_suspicious_commands(sensitivity_text)
    if suspicious:
        warnings.append(
            f"Security: this file {'; '.join(suspicious)}. Sortly only moves files and never runs "
            "them, but review it before you run it yourself."
        )

    return EvidenceBundle(
        path=str(path),
        fingerprint=fingerprint,
        filename=path.name,
        extension=extension,
        parent_directory=path.parent.name,
        content=content,
        metadata=metadata,
        sources=sources,
        warnings=warnings,
        detected_type=detected_type,
        detected_type_confidence=detection_confidence,
        extension_mismatch=(
            detected_type != "unknown"
            and bool(extension)
            and not _extensions_match(extension, detected_type)
        ),
        parser_status=extraction.parser_status,
        extraction_status=status,
        ocr_used=ocr_used,
        sensitive_findings=[finding.to_dict() for finding in findings],
        redacted_content=redact_sensitive(sensitivity_text, [path.name, path.parent.name, *map(str, metadata.values())]),
        suspicious_commands=suspicious,
    )
